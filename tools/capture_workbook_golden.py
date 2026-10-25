"""Capture the workbook's inputs and computed outputs as an anonymised golden file for the engine test.

Usage: python tools/capture_workbook_golden.py WORKBOOK.xlsx OUT.json

The workbook must have been opened and saved in Excel so that its cached values exist (it is read with
``data_only=True``). Location names are replaced by "Site A", "Site B" and so on; no comments or risk text are copied.
Dates are written as ISO strings; the Weekly Tracker rows keep the workbook's Monday report dates.
"""

from __future__ import annotations

import json
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

WEEK_COLUMNS = {
    "cum_target": 5,  # E
    "cum_actual": 6,  # F
    "planned_remaining": 7,  # G
    "actual_remaining": 8,  # H
    "pct_complete": 9,  # I
    "moving_average": 10,  # J
    "forecast_remaining": 11,  # K
}
SETTINGS_OUTPUTS = {
    "forecast_based_on": "B18",
    "forecast_rate": "B19",
    "remaining_at_basis": "B20",
    "planned_completion": "B21",
    "forecast_completion": "B22",
}


def _blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _date(value) -> str | None:
    if _blank(value):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        return from_excel(value).date().isoformat()
    raise SystemExit(f"Expected a date, found {value!r}")


def _number(value):
    if _blank(value):
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise SystemExit(f"Expected a number, found {value!r}")
    return int(value) if float(value).is_integer() else float(value)


def capture(path: str) -> dict:
    wb = load_workbook(path, read_only=False, data_only=True)
    settings, groups_sheet, weekly = wb["Settings"], wb["By Location"], wb["Weekly Tracker"]
    if _blank(settings["B13"].value):
        raise SystemExit("No cached values: open the workbook in Excel, save it, and run this script again.")

    groups = []
    for row in range(5, 11):
        name = groups_sheet.cell(row, 1).value
        if _blank(name) or str(name).strip().lower().startswith("total"):
            if not _blank(name):
                break
            continue
        groups.append(
            {
                "name": f"Site {chr(ord('A') + len(groups))}",
                "baseline": _number(groups_sheet.cell(row, 2).value) or 0,
                "starting_completed": _number(groups_sheet.cell(row, 3).value) or 0,
                "ready_not_complete": _number(groups_sheet.cell(row, 4).value) or 0,
            }
        )

    weeks = []
    for row in range(5, 57):
        report_date = _date(weekly.cell(row, 1).value)
        if report_date is None:
            continue
        weeks.append(
            {
                "report_date": report_date,
                "target": _number(weekly.cell(row, 2).value),
                "actual": _number(weekly.cell(row, 3).value),
                "expected": {key: _number(weekly.cell(row, col).value) for key, col in WEEK_COLUMNS.items()},
            }
        )

    expected = {}
    for key, cell in SETTINGS_OUTPUTS.items():
        value = settings[cell].value
        expected[key] = _date(value) if key.endswith(("_on", "_completion")) else _number(value)
    return {
        "settings": {
            "first_report_date": _date(settings["B13"].value),
            "last_report_date": _date(settings["B14"].value),
            "target_completion": _date(settings["B8"].value),
            "completed_before_first": _number(settings["B9"].value) or 0,
            "moving_average_weeks": _number(settings["B10"].value) or 3,
        },
        "groups": groups,
        "weeks": weeks,
        "expected": expected,
    }


def main(argv: list[str]) -> int:
    if len(argv) != 3:
        print(__doc__)
        return 2
    golden = capture(argv[1])
    Path(argv[2]).write_text(json.dumps(golden, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {argv[2]}: {len(golden['groups'])} groups, {len(golden['weeks'])} weeks.")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
