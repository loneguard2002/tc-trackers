"""One-time import of the TC Windows 11 tracking workbook (spec section 9).

The workbook is untrusted input. It is opened with ``data_only=False`` so that no cached value is ever relied upon;
only input cells are read, and the few that may hold formulas are evaluated by ``eval_simple_formula``, which accepts
nothing beyond arithmetic on numeric constants.
"""

from __future__ import annotations

import ast
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from fractions import Fraction
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from sqlalchemy.orm import Session

from ..audit import record
from ..engine.weeks import max_last_week_ending, report_date_to_week_ending, week_endings
from ..models import Group, ReportInput, Risk, Tracker, TrackerRole, Week

SHEET_SETTINGS = "Settings"
SHEET_GROUPS = "By Location"
SHEET_WEEKS = "Weekly Tracker"
SHEET_RISKS = "Risks & Issues"
SHEET_INPUTS = "Report Inputs"

GROUP_ROWS = range(5, 11)
WEEK_ROWS = range(5, 57)
RISK_ROWS = range(5, 20)
RAG_VALUES = ("Green", "Amber", "Red")
RISK_STATUSES = ("Open", "In progress", "Closed")
REPORT_DAYS = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday")
MAX_FORMULA_LENGTH = 200
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024

_FORMULA_RE = re.compile(r"^=\s*[\d\s+\-*/().]+$")


class ImportError_(Exception):  # noqa: N818  (the trailing underscore avoids shadowing the builtin ImportError)
    """The workbook cannot be imported; the message names the cell or sheet at fault."""


class ImportMismatch(Exception):
    """The group starting figures do not reconcile with the weekly history."""

    def __init__(self, mismatch: int):
        super().__init__(
            f"The group starting figures differ from the quantity completed before the first report plus the weekly "
            f"actuals by {mismatch}."
        )
        self.mismatch = mismatch


@dataclass
class ImportPreview:
    settings: dict = field(default_factory=dict)
    groups: list[dict] = field(default_factory=list)
    weeks: list[dict] = field(default_factory=list)
    risks: list[dict] = field(default_factory=list)
    inputs: dict = field(default_factory=dict)
    group_total: int = 0
    weekly_total: int = 0
    mismatch: int = 0

    def to_json(self) -> dict:
        """A JSON-serialisable form; dates become ``{"$date": "YYYY-MM-DD"}``."""
        return _encode(
            {
                "settings": self.settings,
                "groups": self.groups,
                "weeks": self.weeks,
                "risks": self.risks,
                "inputs": self.inputs,
                "group_total": self.group_total,
                "weekly_total": self.weekly_total,
                "mismatch": self.mismatch,
            }
        )

    @classmethod
    def from_json(cls, data: dict) -> ImportPreview:
        try:
            decoded = _decode(data)
            return cls(
                settings=dict(decoded["settings"]),
                groups=list(decoded["groups"]),
                weeks=list(decoded["weeks"]),
                risks=list(decoded["risks"]),
                inputs=dict(decoded["inputs"]),
                group_total=int(decoded["group_total"]),
                weekly_total=int(decoded["weekly_total"]),
                mismatch=int(decoded["mismatch"]),
            )
        except (KeyError, TypeError, ValueError) as exc:
            raise ImportError_("The import preview is malformed.") from exc


def _encode(value):
    if isinstance(value, date):
        return {"$date": value.isoformat()}
    if isinstance(value, dict):
        return {k: _encode(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_encode(v) for v in value]
    return value


def _decode(value):
    if isinstance(value, dict):
        if set(value) == {"$date"}:
            return date.fromisoformat(value["$date"])
        return {k: _decode(v) for k, v in value.items()}
    if isinstance(value, list):
        return [_decode(v) for v in value]
    return value


# --- Cell evaluation -----------------------------------------------------------------------------------------


def _whole(value: Fraction, cell: str) -> int:
    if value.denominator != 1:
        raise ImportError_(f"{cell} does not evaluate to a whole number.")
    return int(value)


def _eval_node(node: ast.AST, cell: str) -> Fraction:
    if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)) and not isinstance(node.value, bool):
        return Fraction(str(node.value))
    if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
        operand = _eval_node(node.operand, cell)
        return operand if isinstance(node.op, ast.UAdd) else -operand
    if isinstance(node, ast.BinOp) and isinstance(node.op, (ast.Add, ast.Sub, ast.Mult, ast.Div)):
        left = _eval_node(node.left, cell)
        right = _eval_node(node.right, cell)
        if isinstance(node.op, ast.Add):
            return left + right
        if isinstance(node.op, ast.Sub):
            return left - right
        if isinstance(node.op, ast.Mult):
            return left * right
        if right == 0:
            raise ImportError_(f"{cell} divides by zero.")
        return left / right
    raise ImportError_(f"{cell} holds a formula that is not simple arithmetic.")


def eval_simple_formula(value, cell: str = "cell") -> int | None:
    """Return the whole number held by an input cell, or ``None`` when it is blank.

    Accepts integers, whole-valued floats and strings such as ``=202+135``. Anything else raises ``ImportError_``.
    """
    if value is None:
        return None
    if isinstance(value, bool):
        raise ImportError_(f"{cell} holds a value that is not a number.")
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        try:
            return _whole(Fraction(str(value)), cell)
        except ValueError as exc:  # nan and infinity
            raise ImportError_(f"{cell} holds a value that is not a number.") from exc
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        if len(text) > MAX_FORMULA_LENGTH or not _FORMULA_RE.match(text):
            raise ImportError_(f"{cell} holds a formula that is not simple arithmetic.")
        try:
            tree = ast.parse(text[1:].strip(), mode="eval")
            return _whole(_eval_node(tree.body, cell), cell)
        except (SyntaxError, ValueError, RecursionError, MemoryError) as exc:
            raise ImportError_(f"{cell} holds a formula that cannot be read.") from exc
    raise ImportError_(f"{cell} holds a value that is not a number.")


def to_date(value, cell: str) -> date | None:
    """Normalise a datetime, date or Excel serial number to a ``date``; blank cells give ``None``."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            converted = from_excel(value)
        except (ValueError, OverflowError, TypeError) as exc:
            raise ImportError_(f"{cell} holds a number that is not a valid date.") from exc
        if isinstance(converted, datetime):
            return converted.date()
        if isinstance(converted, date):
            return converted
    raise ImportError_(f"{cell} does not hold a date.")


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _count(value, cell: str, default: int | None = None) -> int:
    number = eval_simple_formula(value, cell)
    if number is None:
        if default is None:
            raise ImportError_(f"{cell} is empty.")
        return default
    return number


def _non_negative(value, cell: str) -> int:
    number = _count(value, cell, default=0)
    if number < 0:
        raise ImportError_(f"{cell} must not be negative.")
    return number


def _rag(value, cell: str) -> str | None:
    text = _text(value)
    if not text:
        return None
    for allowed in RAG_VALUES:
        if text.lower() == allowed.lower():
            return allowed
    raise ImportError_(f"{cell} must be Green, Amber or Red.")


# --- Reading the workbook ------------------------------------------------------------------------------------


def _open(path_or_file):
    if isinstance(path_or_file, (str, Path)):
        data = Path(path_or_file).read_bytes()
    elif hasattr(path_or_file, "read"):
        data = path_or_file.read()
    else:
        data = bytes(path_or_file)
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise ImportError_("The workbook is too large to import.")
        return load_workbook(BytesIO(data), read_only=False, data_only=False)
    except ImportError_:
        raise
    except Exception as exc:  # untrusted input: any parser failure means "not a readable workbook"
        raise ImportError_("The file could not be read as an Excel workbook (.xlsx).") from exc


def _sheet(workbook, name: str):
    if name not in workbook.sheetnames:
        raise ImportError_(f'The workbook has no sheet named "{name}".')
    return workbook[name]


def _ref(sheet, coordinate: str) -> str:
    return f"{sheet.title}!{coordinate}"


def _snap_back(report_date: date, report_day: int) -> date:
    """Move a date back to ``report_day``, as the workbook's calculated report dates do."""
    return report_date - timedelta(days=(report_date.weekday() - report_day) % 7)


def _read_settings(workbook, weekday: int):
    sheet = _sheet(workbook, SHEET_SETTINGS)
    day_text = _text(sheet["B5"].value) or "Monday"
    matches = [i for i, name in enumerate(REPORT_DAYS) if name.lower() == day_text.lower()]
    if not matches:
        raise ImportError_(f"{_ref(sheet, 'B5')} must be a weekday from Monday to Friday.")
    report_day = matches[0]

    first_raw = to_date(sheet["B6"].value, _ref(sheet, "B6"))
    last_raw = to_date(sheet["B7"].value, _ref(sheet, "B7"))
    if first_raw is None:
        raise ImportError_(f"{_ref(sheet, 'B6')} is empty.")
    if last_raw is None:
        raise ImportError_(f"{_ref(sheet, 'B7')} is empty.")
    target = to_date(sheet["B8"].value, _ref(sheet, "B8"))
    before_first = _non_negative(sheet["B9"].value, _ref(sheet, "B9"))
    ma_weeks = _count(sheet["B10"].value, _ref(sheet, "B10"), default=3)
    if not 1 <= ma_weeks <= 12:
        raise ImportError_(f"{_ref(sheet, 'B10')} must be between 1 and 12.")

    first_report = _snap_back(first_raw, report_day)
    # The workbook's "last report date used": whole weeks after the first, at least the first, at most 51 weeks.
    whole_weeks = (last_raw - first_report).days // 7
    last_report = first_report + timedelta(weeks=min(51, max(0, whole_weeks)))
    first_we = report_date_to_week_ending(first_report, weekday)
    last_we = min(report_date_to_week_ending(last_report, weekday), max_last_week_ending(first_we))
    settings = {
        "weekday": weekday,
        "first_week_ending": first_we,
        "last_week_ending": last_we,
        "target_completion": target,
        "completed_before_first": before_first,
        "moving_average_weeks": ma_weeks,
    }
    return settings, first_report


def _read_groups(workbook) -> list[dict]:
    sheet = _sheet(workbook, SHEET_GROUPS)
    groups: list[dict] = []
    seen: set[str] = set()
    for row in GROUP_ROWS:
        name = _text(sheet.cell(row, 1).value)
        if name.lower().startswith("total"):
            break
        if not name:
            continue
        if name in seen:
            raise ImportError_(f"{_ref(sheet, f'A{row}')} repeats the name of another row.")
        seen.add(name)
        groups.append(
            {
                "name": name,
                "baseline": _non_negative(sheet.cell(row, 2).value, _ref(sheet, f"B{row}")),
                "starting_completed": _non_negative(sheet.cell(row, 3).value, _ref(sheet, f"C{row}")),
                "ready_not_complete": _non_negative(sheet.cell(row, 4).value, _ref(sheet, f"D{row}")),
                "planned_completion": to_date(sheet.cell(row, 7).value, _ref(sheet, f"G{row}")),
                "rag": _rag(sheet.cell(row, 8).value, _ref(sheet, f"H{row}")),
                "comment": _text(sheet.cell(row, 9).value),
            }
        )
    return groups


def _read_weeks(workbook, first_report: date, last_report_count: int, weekday: int) -> list[dict]:
    sheet = _sheet(workbook, SHEET_WEEKS)
    weeks: list[dict] = []
    for index, row in enumerate(WEEK_ROWS):
        if index >= last_report_count:
            break
        report_date = first_report + timedelta(weeks=index)
        target = eval_simple_formula(sheet.cell(row, 2).value, _ref(sheet, f"B{row}"))
        actual = eval_simple_formula(sheet.cell(row, 3).value, _ref(sheet, f"C{row}"))
        if actual is not None and actual < 0:
            raise ImportError_(f"{_ref(sheet, f'C{row}')} must not be negative.")
        weeks.append(
            {"week_ending": report_date_to_week_ending(report_date, weekday), "target": target, "actual": actual}
        )
    return weeks


def _read_risks(workbook) -> list[dict]:
    sheet = _sheet(workbook, SHEET_RISKS)
    risks: list[dict] = []
    seen: set[str] = set()
    for row in RISK_ROWS:
        description = _text(sheet.cell(row, 2).value)
        if not description:
            continue
        ref = _text(sheet.cell(row, 1).value) or f"R{row - 4}"
        if ref in seen:
            raise ImportError_(f"{_ref(sheet, f'A{row}')} repeats the ID of another row.")
        seen.add(ref)
        status = _text(sheet.cell(row, 8).value) or "Open"
        matched = [s for s in RISK_STATUSES if s.lower() == status.lower()]
        if not matched:
            raise ImportError_(f"{_ref(sheet, f'H{row}')} must be Open, In progress or Closed.")
        risks.append(
            {
                "ref": ref,
                "description": description,
                "impact": _text(sheet.cell(row, 3).value),
                "group": _text(sheet.cell(row, 4).value),
                "owner_role": _text(sheet.cell(row, 5).value),
                "raised_on": to_date(sheet.cell(row, 6).value, _ref(sheet, f"F{row}")),
                "target_date": to_date(sheet.cell(row, 7).value, _ref(sheet, f"G{row}")),
                "status": matched[0],
                "rag": _rag(sheet.cell(row, 9).value, _ref(sheet, f"I{row}")),
                "latest_update": _text(sheet.cell(row, 10).value),
            }
        )
    return risks


def _read_inputs(workbook, weekday: int) -> dict:
    sheet = _sheet(workbook, SHEET_INPUTS)
    report_date = to_date(sheet["B5"].value, _ref(sheet, "B5"))
    status = _rag(sheet["B7"].value, _ref(sheet, "B7"))
    headline = _text(sheet["B8"].value)
    decisions = _text(sheet["B9"].value)
    if report_date is None or not (status or headline or decisions):
        return {}
    return {
        "week_ending": report_date_to_week_ending(report_date, weekday),
        "status": status,
        "headline": headline,
        "decisions": decisions,
    }


def read_workbook(path_or_file, weekday: int) -> ImportPreview:
    """Read the workbook's input cells into a preview; raises ``ImportError_`` for anything unusable."""
    if not isinstance(weekday, int) or isinstance(weekday, bool) or not 0 <= weekday <= 6:
        raise ImportError_("The week-ending day must be 0 (Monday) to 6 (Sunday).")
    workbook = _open(path_or_file)
    try:
        settings, first_report = _read_settings(workbook, weekday)
        count = len(week_endings(settings["first_week_ending"], settings["last_week_ending"]))
        groups = _read_groups(workbook)
        if not groups:
            raise ImportError_(f"{SHEET_GROUPS} lists no groups in rows 5 to 10.")
        weeks = _read_weeks(workbook, first_report, count, weekday)
        risks = _read_risks(workbook)
        inputs = _read_inputs(workbook, weekday)
    finally:
        workbook.close()
    group_total = sum(g["starting_completed"] for g in groups)
    weekly_total = sum(w["actual"] for w in weeks if w["actual"] is not None)
    mismatch = group_total - (settings["completed_before_first"] + weekly_total)
    return ImportPreview(settings, groups, weeks, risks, inputs, group_total, weekly_total, mismatch)


# --- Creating the tracker ------------------------------------------------------------------------------------


def create_tracker_from_preview(
    session: Session,
    preview: ImportPreview,
    name: str,
    unit_label: str,
    group_label: str,
    user_id: int,
    accept_mismatch: bool,
) -> Tracker:
    """Add the tracker and its history to ``session`` (the caller commits)."""
    name = name.strip()
    if not name:
        raise ImportError_("A tracker name is required.")
    if preview.mismatch != 0 and not accept_mismatch:
        raise ImportMismatch(preview.mismatch)
    s = preview.settings
    try:
        week_endings(s["first_week_ending"], s["last_week_ending"])
    except (KeyError, ValueError) as exc:
        raise ImportError_(f"The imported week range is not valid: {exc}") from exc
    if not 1 <= s["moving_average_weeks"] <= 12:
        raise ImportError_("The moving-average period must be between 1 and 12.")

    tracker = Tracker(
        name=name,
        unit_label=unit_label.strip(),
        group_label=group_label.strip(),
        week_ending_weekday=s["weekday"],
        first_week_ending=s["first_week_ending"],
        last_week_ending=s["last_week_ending"],
        target_completion=s["target_completion"],
        completed_before_first=s["completed_before_first"],
        moving_average_weeks=s["moving_average_weeks"],
    )
    session.add(tracker)
    session.flush()

    group_ids: dict[str, int] = {}
    for order, item in enumerate(preview.groups):
        group = Group(tracker_id=tracker.id, sort_order=order, **item)
        session.add(group)
        session.flush()
        group_ids[group.name] = group.id
    for item in preview.weeks:
        session.add(
            Week(
                tracker_id=tracker.id,
                week_ending=item["week_ending"],
                target=item["target"],
                imported_actual=item["actual"],
            )
        )
    for item in preview.risks:
        values = {k: v for k, v in item.items() if k != "group"}
        session.add(Risk(tracker_id=tracker.id, group_id=group_ids.get(item["group"]), **values))
    if preview.inputs:
        session.add(ReportInput(tracker_id=tracker.id, prepared_by=user_id, **preview.inputs))
    session.add(TrackerRole(user_id=user_id, tracker_id=tracker.id, role="owner"))
    record(
        session,
        user_id=user_id,
        tracker_id=tracker.id,
        action="tracker_imported",
        entity="tracker",
        entity_id=tracker.id,
        detail={
            "name": name,
            "groups": len(preview.groups),
            "weeks": len(preview.weeks),
            "risks": len(preview.risks),
            "mismatch": preview.mismatch,
            "mismatch_accepted": preview.mismatch != 0,
        },
    )
    session.flush()
    return tracker
