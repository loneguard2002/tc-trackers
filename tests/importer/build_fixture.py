"""Build an anonymised workbook with the same sheet names and input cells as the real tracking workbook."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook

TARGETS = [35, 40, 40, 40]
ACTUALS = [40, 38, 20, 40]  # sum 138 = 135 + 0 + 3


def build_workbook(*, deployed_site_a: int = 135, baseline_site_a="=202+135", tweak=None) -> Workbook:
    wb = Workbook()
    wb.active.title = "Dashboard"
    inputs = wb.create_sheet("Report Inputs")
    groups = wb.create_sheet("By Location")
    weekly = wb.create_sheet("Weekly Tracker")
    risks = wb.create_sheet("Risks & Issues")
    settings = wb.create_sheet("Settings")

    settings["B5"] = "Monday"
    settings["B6"] = datetime(2026, 9, 7)
    settings["B7"] = datetime(2026, 9, 28)
    settings["B8"] = datetime(2026, 12, 18)
    settings["B9"] = 0
    settings["B10"] = 3
    settings["B13"] = '=IF(B6="","",B6)'  # calculated cells are never read

    for row, values in enumerate(
        [("Site A", baseline_site_a, deployed_site_a, 5), ("Site B", 2, 0, 0), ("Remote", 30, 3, 1)], start=5
    ):
        groups.cell(row, 1, values[0])
        groups.cell(row, 2, values[1])
        groups.cell(row, 3, values[2])
        groups.cell(row, 4, values[3])
    groups["G5"] = datetime(2026, 11, 20)
    groups["H5"] = "Amber"
    groups["I5"] = "Awaiting delivery"
    groups["A8"] = "Total"
    groups["B8"] = "=SUM(B5:B7)"

    weekly["A5"] = "=Settings!$B$13"
    for index, (target, actual) in enumerate(zip(TARGETS, ACTUALS, strict=True)):
        weekly.cell(5 + index, 2, target)
        weekly.cell(5 + index, 3, actual)

    risks["A5"], risks["B5"], risks["C5"], risks["D5"] = "R1", "Delivery delay", "Slips the plan", "Site A"
    risks["E5"], risks["F5"], risks["G5"] = "Project manager", datetime(2026, 9, 1), datetime(2026, 10, 1)
    risks["H5"], risks["I5"], risks["J5"] = "Open", "Amber", "Supplier chased"

    inputs["B5"] = datetime(2026, 9, 28)
    inputs["B7"] = "Green"
    inputs["B8"] = "Deployment is on plan."
    inputs["B9"] = "None required."

    if tweak:
        tweak(wb)
    return wb


def workbook_bytes(**kwargs) -> bytes:
    buffer = BytesIO()
    build_workbook(**kwargs).save(buffer)
    return buffer.getvalue()
